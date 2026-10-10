"""Convierte Catalogo_La_Plaza.xlsx en los datos de la tienda.

Genera:
  data/catalogo.json         perfumes, cupones y configuración
  producto/<perfume>.html    una página por perfume (para Google y para compartir)
  categoria/<nombre>.html    páginas de categoría
  sitemap.xml y robots.txt   solo si cargaste «URL del sitio» en Configuración

Uso (desde la carpeta de la tienda):
    pip install openpyxl
    python herramientas/excel_a_json.py Catalogo_La_Plaza.xlsx
"""
import html
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DECANT_SIZES = (3, 5, 10)


def slug(text):
    t = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "-", t.lower()).strip("-")


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    digits = re.sub(r"[^\d]", "", str(v))
    return int(digits) if digits else None


def text(v):
    return str(v).strip() if v not in (None, "") else ""


def si(v):
    return text(v).lower() in ("sí", "si", "s", "yes")


def lista(v):
    return [p.strip()[0].upper() + p.strip()[1:] for p in text(v).split(",") if p.strip()]


def leer_productos(ws):
    headers = [text(c.value) for c in ws[1]]
    col = {h: i for i, h in enumerate(headers)}

    def get(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    productos, usados = [], set()
    for orden, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        pid = text(get(row, "ID"))
        nombre = text(get(row, "Nombre"))
        if not pid or not nombre or pid.upper() == "EJEMPLO":
            continue
        if not si(get(row, "Activo")):
            continue
        marca = text(get(row, "Marca"))
        s = slug(f"{marca} {nombre}")
        while s in usados:
            s += "-2"
        usados.add(s)

        oferta = num(get(row, "Oferta %"))
        oferta = oferta if oferta and 0 < oferta < 91 else None

        ml_dec = num(get(row, "ml disponibles para decants")) or 0
        decants = []
        for ml in DECANT_SIZES:
            precio = num(get(row, f"Precio {ml} ml (Gs)"))
            if precio:
                decants.append({"ml": ml, "precio": precio, "disponible": ml_dec >= ml})

        botella = None
        precio_b = num(get(row, "Precio botella (Gs)"))
        if precio_b:
            botella = {
                "precio": precio_b,
                "ml": num(get(row, "Tamaño botella (ml)")),
                # None = sin dato cargado: la tienda muestra «consultar disponibilidad»
                "stock": num(get(row, "Botellas en stock")),
            }

        productos.append({
            "id": s,
            "orden": orden,
            "marca": marca,
            "nombre": nombre,
            "genero": text(get(row, "Género")),
            "tipo": text(get(row, "Tipo")),
            "familia": text(get(row, "Familia olfativa")).capitalize(),
            "concentracion": text(get(row, "Concentración")),
            "notas": {
                "salida": text(get(row, "Notas de salida")),
                "corazon": text(get(row, "Notas de corazón")),
                "fondo": text(get(row, "Notas de fondo")),
            },
            "descripcion": text(get(row, "Descripción corta")),
            "fotos": [f.strip() for f in text(get(row, "Foto (archivo)")).split(",") if f.strip()],
            "destacado": si(get(row, "Destacado")),
            "novedad": si(get(row, "Novedad")),
            "masVendido": si(get(row, "Más vendido")),
            "oferta": oferta,
            "momento": text(get(row, "Momento")),
            "clima": text(get(row, "Clima")),
            "ocasion": lista(get(row, "Ocasión")),
            "intensidad": text(get(row, "Intensidad")),
            "duracion": text(get(row, "Duración")),
            "botella": botella,
            "decants": decants,
            "mlDecant": ml_dec,
        })
    return productos


def leer_cupones(wb):
    if "Cupones" not in wb.sheetnames:
        return []
    out = []
    for row in wb["Cupones"].iter_rows(min_row=2, values_only=True):
        codigo = text(row[0] if row else "").upper()
        if not codigo or codigo == "EJEMPLO10" or not si(row[4] if len(row) > 4 else ""):
            continue
        tipo = "monto" if "monto" in text(row[1]).lower() else "porcentaje"
        valor = num(row[2])
        if not valor:
            continue
        out.append({"codigo": codigo, "tipo": tipo, "valor": valor, "minimo": num(row[3]) or 0})
    return out


def leer_config(ws):
    cfg = {"whatsapp": "", "instagram": "", "tiktok": "", "retiro": "", "pagos": [], "premium": None,
           "url": "", "banner": "", "zonas": []}
    keys = {
        "WhatsApp": "whatsapp", "Usuario de Instagram": "instagram", "Usuario de TikTok": "tiktok",
        "Lugar de retiro": "retiro", "Métodos de pago": "pagos", "Costo frasco premium": "premium",
        "URL del sitio": "url", "Banner promocional": "banner",
    }
    rows = list(ws.iter_rows(values_only=True))
    zonas_desde = None
    for i, row in enumerate(rows):
        label = text(row[0]) if row else ""
        for k, field in keys.items():
            if label.startswith(k):
                v = row[1] if len(row) > 1 else None
                if field == "pagos":
                    cfg[field] = [p.strip() for p in text(v).split(",") if p.strip()]
                elif field == "premium":
                    cfg[field] = num(v)
                elif field in ("whatsapp", "instagram", "tiktok"):
                    cfg[field] = re.sub(r"[^\w.]", "", text(v))
                elif field == "url":
                    u = text(v)
                    cfg[field] = (u if u.endswith("/") else u + "/") if u else ""
                else:
                    cfg[field] = text(v)
        if label.startswith("Zona / ciudad"):
            zonas_desde = i + 1
    if zonas_desde is not None:
        for row in rows[zonas_desde:]:
            zona = text(row[0]) if row else ""
            if not zona:
                continue
            raw = row[1] if len(row) > 1 else None
            costo = num(raw) if isinstance(raw, (int, float)) or (raw and re.search(r"\d", str(raw))) else None
            txt = text(raw).lower()
            gratis = "sin costo" in txt or "gratis" in txt
            cfg["zonas"].append({
                "zona": zona[0].upper() + zona[1:],
                "costo": 0 if gratis else costo,
                "consultar": costo is None and not gratis,
                "tiempo": text(row[2]) if len(row) > 2 else "",
            })
    return cfg


def gs(n):
    return "Gs. " + f"{n:,}".replace(",", ".")


def pagina(template, ruta, titulo, descripcion, imagen, cfg, jsonld=None, canonical=""):
    """Copia de index.html con su propio título y descripción, que abre la ruta indicada."""
    e = html.escape
    t = re.sub(r"<title>.*?</title>", f"<title>{e(titulo)}</title>", template, flags=re.S)
    t = re.sub(r'(<meta name="description" content=")[^"]*', lambda m: m.group(1) + e(descripcion), t)
    t = re.sub(r'(<meta property="og:title" content=")[^"]*', lambda m: m.group(1) + e(titulo), t)
    t = re.sub(r'(<meta property="og:description" content=")[^"]*', lambda m: m.group(1) + e(descripcion), t)
    extra = '<base href="../">\n'
    if imagen:
        img = (cfg["url"] + imagen) if cfg["url"] else imagen
        extra += f'<meta property="og:image" content="{e(img)}">\n'
    if canonical:
        extra += f'<link rel="canonical" href="{e(canonical)}">\n'
    if jsonld:
        extra += '<script type="application/ld+json">' + json.dumps(jsonld, ensure_ascii=False).replace("</", "<\\/") + "</script>\n"
    extra += f"<script>window.INITIAL_ROUTE = {json.dumps(ruta)};</script>\n"
    return t.replace("<meta charset=\"utf-8\">", "<meta charset=\"utf-8\">\n" + extra, 1)


def precio_desde(p):
    precios = []
    if p["botella"]:
        precios.append(p["botella"]["precio"])
    precios += [d["precio"] for d in p["decants"] if d["disponible"]]
    base = min(precios) if precios else 0
    return round(base * (100 - p["oferta"]) / 100 / 100) * 100 if p["oferta"] else base


def generar_paginas(data):
    cfg = data["config"]
    template = (ROOT / "index.html").read_text(encoding="utf-8")
    urls = [cfg["url"]] if cfg["url"] else []
    pdir = ROOT / "producto"
    pdir.mkdir(exist_ok=True)
    for old in pdir.glob("*.html"):
        old.unlink()
    for p in data["productos"]:
        titulo = f"{p['marca']} {p['nombre']} — La Plaza Perfumería"
        desc = (p["descripcion"] + " " if p["descripcion"] else "") + f"Desde {gs(precio_desde(p))}. {cfg['retiro']}, Paraguay."
        foto = f"img/{p['fotos'][0]}" if p["fotos"] else ""
        canon = f"{cfg['url']}producto/{p['id']}.html" if cfg["url"] else ""
        ld = {"@context": "https://schema.org", "@type": "Product", "name": p["nombre"],
              "brand": {"@type": "Brand", "name": p["marca"]}, "description": p["descripcion"],
              "offers": {"@type": "AggregateOffer", "priceCurrency": "PYG", "lowPrice": precio_desde(p)}}
        if foto:
            ld["image"] = (cfg["url"] + foto) if cfg["url"] else foto
        (pdir / f"{p['id']}.html").write_text(
            pagina(template, f"/p/{p['id']}", titulo, desc, foto, cfg, ld, canon), encoding="utf-8")
        if canon:
            urls.append(canon)

    cdir = ROOT / "categoria"
    cdir.mkdir(exist_ok=True)
    for old in cdir.glob("*.html"):
        old.unlink()
    cats = [
        ("masculinos", "Perfumes masculinos", "/catalogo?genero=Masculino", lambda p: p["genero"] == "Masculino"),
        ("femeninos", "Perfumes femeninos", "/catalogo?genero=Femenino", lambda p: p["genero"] == "Femenino"),
        ("unisex", "Perfumes unisex", "/catalogo?genero=Unisex", lambda p: p["genero"] == "Unisex"),
        ("decants", "Decants de perfumes", "/decants", lambda p: any(d["disponible"] for d in p["decants"])),
    ]
    for key, nombre, ruta, test in cats:
        items = [p for p in data["productos"] if test(p)]
        if not items:
            continue
        marcas = sorted({p["marca"] for p in items})
        desc = f"{nombre} originales en Paraguay: {', '.join(marcas[:6])}. {len(items)} fragancias, retiro en {cfg['retiro']}."
        canon = f"{cfg['url']}categoria/{key}.html" if cfg["url"] else ""
        (cdir / f"{key}.html").write_text(
            pagina(template, ruta, f"{nombre} — La Plaza Perfumería", desc, "", cfg, None, canon), encoding="utf-8")
        if canon:
            urls.append(canon)

    if cfg["url"]:
        hoy = date.today().isoformat()
        body = "".join(f"<url><loc>{html.escape(u)}</loc><lastmod>{hoy}</lastmod></url>" for u in urls)
        (ROOT / "sitemap.xml").write_text(
            f'<?xml version="1.0" encoding="UTF-8"?>\n<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">{body}</urlset>\n',
            encoding="utf-8")
        (ROOT / "robots.txt").write_text(f"User-agent: *\nAllow: /\nSitemap: {cfg['url']}sitemap.xml\n", encoding="utf-8")
        return len(urls)
    return 0


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "Catalogo_La_Plaza.xlsx")
    wb = load_workbook(src, data_only=True)
    data = {
        "config": leer_config(wb["Configuración"]),
        "productos": leer_productos(wb["Productos"]),
        "cupones": leer_cupones(wb),
    }
    out = ROOT / "data" / "catalogo.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    n = generar_paginas(data)
    print(f"{len(data['productos'])} productos, {len(data['cupones'])} cupones activos → {out}")
    print(f"Páginas generadas en producto/ y categoria/." + (f" Sitemap con {n} direcciones." if n else
          " Sin sitemap: cargá «URL del sitio» en Configuración."))


if __name__ == "__main__":
    main()
